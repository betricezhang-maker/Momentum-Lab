"""Portable Excel review exchange. Uploaded cells are decisions, never authority."""
import base64
import copy
import hashlib
import io
import json
import uuid
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from .data_integrity import DATA_LOCK, save_acknowledgement, revoke_acknowledgement
from .serialization import dumps

EDITABLE = ['Requested action', 'User reason', 'Evidence reference or URL']
PROTECTED = ['Review group ID', 'Issue IDs', 'Dataset identity', 'Dataset fingerprint', 'Finding fingerprint',
             'Ticker', 'ETF name', 'Affected dates', 'Component', 'Original classification', 'Original severity',
             'Current status', 'Blocks research', 'Existing explanation', 'Existing evidence', 'Related raw/adjusted IDs']
COLUMNS = PROTECTED + EDITABLE
# Display/provenance columns may be reformatted or refreshed by Excel. These
# fields are retained in the workbook but are not used as the authorization key.
IDENTITY_FIELDS = ['Review group ID', 'Issue IDs', 'Dataset identity',
                   'Finding fingerprint', 'Ticker']


def identity_value(field, value):
    text=str(value or '').strip()
    if field=='Issue IDs':
        try:return tuple(sorted(json.loads(text)))
        except (ValueError,TypeError):return text
    return text

def digest(value): return hashlib.sha256(dumps(value,sort_keys=True).encode()).hexdigest()
def folder(paths): return Path(paths['root'])/'integrity_reviews'
def read_json(path): return json.loads(path.read_text(encoding='utf-8'))
def write_json(path, value):
    path.parent.mkdir(parents=True,exist_ok=True)
    temp=path.with_suffix('.'+uuid.uuid4().hex+'.tmp')
    try: temp.write_text(dumps(value,indent=2),encoding='utf-8');temp.replace(path)
    finally: temp.unlink(missing_ok=True)
def safe_id(value):
    if not isinstance(value,str) or len(value)!=32 or any(c not in '0123456789abcdef' for c in value): raise ValueError('Invalid review identifier.')
    return value

def all_findings(report):
    rows=copy.deepcopy(report.get('completeness',{}).get('findings',[]))
    for check in report.get('checks',[]):
        if check['status']=='PASS' or check['name']=='completeness': continue
        rows.append(dict(issue_id='STRUCT-'+digest(check['name'])[:16],ticker='*',date='',end_date='',affected_component=check['name'],
                         classification=check['message'],severity=check['status'],resolution_status='UNRESOLVED',
                         blocks_research=check['status']=='FAIL',overridable=False,evidence_source=dumps(check.get('details',{})),structural=True))
    for row in rows:
        row['scope_label']='Research coverage' if row['affected_component']=='partial_years' else 'Dataset-wide' if row['ticker']=='*' else row['ticker']
        row['repair_eligible']=repair_eligible(row)
        if row['ticker']=='*':
            if row['classification']=='EVIDENCE_LIMITATION':
                row['explanation']=(row.get('signal_impact','Evidence unavailable.')+
                    ' Scope: dataset validation evidence. Missing observations remain detectable and block until individually reviewed; '
                    'this evidence limitation itself does not excuse gaps or establish suspension. No data repair required for this notice.')
            else:
                row['explanation']=row['classification']+' '+row['evidence_source']
            row['action_label']='Inspect and correct the reported condition' if row['blocks_research'] else 'No data repair required'
    return rows

def relevant_signature(row):
    return digest({k:row.get(k) for k in ('ticker','date','end_date','affected_component','classification','severity','session_signatures',
        'evidence_source','resolution_status','explanation','confirmed_at','inherited_from')})

def review_groups(report):
    findings=all_findings(report);used=set();groups=[];by_ticker={}
    for finding in findings:
        if finding['affected_component'] in ('raw_price','adjusted_price'):
            by_ticker.setdefault(finding['ticker'],[]).append(finding)
    for row in sorted(findings,key=lambda r:(r['affected_component']!='raw_price',r['issue_id'])):
        if row['issue_id'] in used: continue
        members=[row]
        if row['affected_component']=='raw_price' and row.get('session_dates'):
            for child in by_ticker.get(row['ticker'],[]):
                if child['issue_id'] in used or child['affected_component']!='adjusted_price' or child['ticker']!=row['ticker']:continue
                if child.get('session_dates')!=row['session_dates'] or not set(row['session_dates']).issubset(child.get('raw_missing_dates',[])):continue
                if child.get('confirmation') and not child.get('inherited_from'):continue
                members.append(child)
        ids=sorted(r['issue_id'] for r in members);used.update(ids)
        related=[r['issue_id'] for r in by_ticker.get(row['ticker'],[])
                 if set(r.get('session_dates',[])) & set(row.get('session_dates',[]))]
        groups.append(dict(group_id='GRP-'+digest(ids)[:20],issue_ids=ids,primary_id=row['issue_id'],members=members,
                           fingerprint=digest([relevant_signature(r) for r in members]),related_ids=related))
    return groups

def export_review(paths, report, names=None, group_ids=None):
    # This is the application's runtime exporter; no Codex/Node runtime is required on the user's PC.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    names=names or {};export_id=uuid.uuid4().hex;groups=review_groups(report);rows={};evidence_rows=[]
    if group_ids is not None:
        selected=set(group_ids)
        groups=[g for g in groups if g['group_id'] in selected]
    workbook=Workbook();sheet=workbook.active;sheet.title='Review'
    sheet.append(COLUMNS);sheet.freeze_panes='C2';sheet.auto_filter.ref=f'A1:S{max(1,len(groups)+1)}'
    for group in groups:
        members=group['members'];first=members[0]
        join=lambda field:'\n'.join(dict.fromkeys(str(r.get(field,'')) for r in members))
        values=[group['group_id'],dumps(group['issue_ids']),report['manifest']['dataset_id'],report['manifest']['dataset_fingerprint'],group['fingerprint'],
                first['ticker'],first['scope_label'] if first['ticker']=='*' else names.get(first['ticker'],'Unavailable'),join('date')+' → '+join('end_date'),join('affected_component'),
                join('classification'),join('severity'),join('resolution_status'),'YES' if any(r.get('blocks_research') for r in members) else 'NO',
                join('explanation'), '\n'.join(dict.fromkeys(str(r.get('confirmation',{}).get('reference') or r.get('evidence_source','')) for r in members)),dumps(group['related_ids'])]
        for index,value in enumerate(values):
            if len(str(value))>32000:
                for offset in range(0,len(str(value)),32000):evidence_rows.append([group['group_id'],PROTECTED[index],str(offset//32000+1),str(value)[offset:offset+32000]])
                values[index]=f'See Evidence sheet: {group["group_id"]} / {PROTECTED[index]}'
        rows[group['group_id']]=dict(protected=dict(zip(PROTECTED,values)),group=group)
        sheet.append(values+['leave unchanged','',''])
    for row in sheet:
        for cell in row:
            cell.data_type='s' # Never interpret source explanations as spreadsheet formulas.
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if cell.row==1:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='16324F')
            elif cell.column>len(PROTECTED):cell.fill=PatternFill('solid',fgColor='FFF2CC');cell.protection=Protection(locked=False)
    widths=[27,35,29,25,25,17,23,27,24,31,18,20,16,42,42,35,22,48,42]
    from openpyxl.utils import get_column_letter
    for index,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(index)].width=width
    sheet.row_dimensions[1].height=42
    for index in range(2,len(groups)+2):sheet.row_dimensions[index].height=60
    choices=DataValidation(type='list',formula1='"leave unchanged,repair,confirm,reopen"');choices.errorTitle='Choose a review action';choices.error='Choose an action from the list.';choices.showErrorMessage=True
    sheet.add_data_validation(choices);choices.add(f'Q2:Q{max(2,len(groups)+1)}')
    meta=workbook.create_sheet('Instructions')
    for row in [('Momentum Lab issue review',export_id),('Template version','2'),('Integrity run ID',report.get('timestamp','')),('Universe',report['universe']),('Exported at',datetime.now(timezone.utc).isoformat()),
                ('Review groups',str(len(groups))),('Underlying findings',str(sum(len(g['members']) for g in groups))),
                ('Edit only yellow columns','Requested action, User reason, Evidence reference or URL.'),
                ('Confirm','Requires a reason. Means manual acceptance, not a verified suspension or repaired data.'),
                ('Linked gaps','One raw/adjusted group shares a reason. Independent confirmations remain separate.'),
                ('Next step','Save as .xlsx, import, inspect preview, then click Apply reviewed decisions.'),
                ('Original fields','Do not change protected identity/evidence fields. Sorting/filtering rows is supported.'),
                ('Repairs','Only missing provider observations are fetched. Local adjusted prices are rebuilt using the existing QFQ method.'),
                ('History','Full original findings and validation evidence are retained in the application export snapshot.')]:meta.append(row)
    meta.column_dimensions['A'].width=27;meta.column_dimensions['B'].width=105
    for row in meta:
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    detail=workbook.create_sheet('Underlying findings');detail.append(['Review group ID','Issue ID','Ticker','Component','Start','End','Original classification','Severity','Status','Blocks research','Explanation','Evidence'])
    detail_rows=[]
    for group in groups:
        for finding in group['members']:
            values=[group['group_id'],finding['issue_id'],finding['ticker'],finding['affected_component'],finding.get('date',''),finding.get('end_date',''),finding['classification'],finding['severity'],finding.get('resolution_status','UNRESOLVED'),'YES' if finding.get('blocks_research') else 'NO',finding.get('explanation',''),finding.get('confirmation',{}).get('reference') or finding.get('evidence_source','')]
            for index,value in enumerate(values):
                if len(str(value))>32000:
                    for offset in range(0,len(str(value)),32000):evidence_rows.append([finding['issue_id'],str(index),str(offset//32000+1),str(value)[offset:offset+32000]])
                    values[index]=f'See Evidence sheet: {finding["issue_id"]} / {index}'
            detail_rows.append(values);detail.append(values)
    detail.freeze_panes='C2';detail.auto_filter.ref=f'A1:L{max(1,len(detail_rows)+1)}'
    for column in range(1,13):detail.column_dimensions[get_column_letter(column)].width=30 if column<11 else 60
    for row in detail:
        for cell in row:cell.data_type='s';cell.alignment=Alignment(wrap_text=True,vertical='top')
    if evidence_rows:
        evidence=workbook.create_sheet('Evidence');evidence.append(['Source ID','Field','Part','Text'])
        for values in evidence_rows:evidence.append(values)
        for row in evidence:
            for cell in row:cell.data_type='s'
        evidence.column_dimensions['D'].width=100
    buffer=io.BytesIO();workbook.save(buffer)
    write_json(folder(paths)/f'export-{export_id}.json',dict(export_id=export_id,universe=report['universe'],rows=rows,report=report,detail_rows=detail_rows,evidence_rows=evidence_rows))
    return dict(ok=True,export_id=export_id,filename=f'MomentumLab_issues_{export_id[:8]}.xlsx',workbook=base64.b64encode(buffer.getvalue()).decode(),groups=len(groups),findings=sum(len(g['members']) for g in groups))

def parse_workbook(paths, encoded, progress=lambda **kw:None):
    from openpyxl import load_workbook
    try: data=base64.b64decode(encoded,validate=True)
    except Exception as exc: raise ValueError('Upload a valid .xlsx workbook.') from exc
    if len(data)>12_000_000:raise ValueError('Workbook exceeds the 12 MB upload limit.')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(i.file_size for i in archive.infolist())>60_000_000:raise ValueError('Expanded workbook is too large.')
    workbook=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
    try:
        export_id=safe_id(workbook['Instructions']['B1'].value)
        snapshot=read_json(folder(paths)/f'export-{export_id}.json')
        for sheet_name,expected in [('Underlying findings',snapshot.get('detail_rows',[])),('Evidence',snapshot.get('evidence_rows',[]))]:
            if not expected:continue
            actual=[]
            for cells in workbook[sheet_name].iter_rows(min_row=2):
                if any(c.data_type=='f' for c in cells):raise ValueError('Protected evidence sheet contains formulas.')
                values=[str(c.value) if c.value is not None else '' for c in cells]
                if any(values):actual.append(values)
            if sorted(actual)!=sorted(expected):raise ValueError(f'Protected {sheet_name} fields were altered.')
        sheet=workbook['Review']
        if sheet.max_row>100_001:raise ValueError('Too many worksheet rows.')
        source=sheet.iter_rows();headers=[c.value for c in next(source)]
        if headers!=COLUMNS:raise ValueError('Column headers were altered. Export a new workbook and edit only the three decision columns.')
        rows=[];total=max(0,sheet.max_row-1)
        progress(stage='PARSING',current=0,total=total,message='Reading reviewed workbook rows.')
        for cells in source:
            values=[str(c.value) if c.value is not None else '' for c in cells]
            if not any(values):continue
            rows.append(dict(values=dict(zip(COLUMNS,values)),formula=any(c.data_type=='f' for c in cells)))
            if len(rows)%50==0:progress(stage='PARSING',current=len(rows),total=total,message='Reading reviewed workbook rows.')
        progress(stage='PARSING',current=len(rows),total=total,message='Workbook rows read.')
        return snapshot,rows
    finally:workbook.close()

def evaluate(snapshot, rows, report, applied):
    current={g['group_id']:g for g in review_groups(report)};result=[];duplicates={}
    for index,input_row in enumerate(rows,2):
        cells=input_row['values'];gid=cells.get('Review group ID');action=cells.get('Requested action','').strip().lower() or 'leave unchanged'
        reason=cells.get('User reason','').strip();reference=cells.get('Evidence reference or URL','').strip()
        item=dict(row=index,group_id=gid,action=action,reason=reason,reference=reference,status='VALID',issue_ids=[],inherited=[])
        exported=snapshot['rows'].get(gid);error=None
        if not exported:error='Unknown review group / issue ID.'
        elif input_row.get('formula'):error='Formula cells are not permitted in reviewed decisions.'
        elif any(identity_value(k,cells.get(k,''))!=identity_value(k,exported['protected'].get(k,'')) for k in IDENTITY_FIELDS):
            changed=[k for k in IDENTITY_FIELDS if identity_value(k,cells.get(k,''))!=identity_value(k,exported['protected'].get(k,''))]
            item['mismatches']=[dict(field=k,expected=exported['protected'].get(k,''),imported=cells.get(k,'')) for k in changed]
            error='Target identity fields were altered: '+', '.join(changed)+'. Re-export a fresh workbook and edit only the yellow decision columns.'
        elif action not in ('leave unchanged','repair','confirm','reopen'):error='Unknown requested action.'
        else:
            group=exported['group'];item.update(issue_ids=group['issue_ids'],primary_id=group['primary_id'])
            item['decision_id']=digest([snapshot['universe'],gid,group['fingerprint'],action,reason,reference])
            actual=current.get(gid)
            if item['decision_id'] in applied:item['status']='ALREADY_APPLIED'
            elif action=='leave unchanged':item['status']='UNCHANGED'
            elif not actual or actual['fingerprint']!=group['fingerprint']:error='Stale finding: relevant evidence or resolution changed. Re-export and review.'
            elif action=='confirm' and not reason:error='Confirmation requires a non-empty user reason.'
            elif action=='confirm' and not actual['members'][0].get('overridable'):error='Structural/evidence-only finding cannot be confirmed.'
            elif action=='reopen' and not any(m.get('confirmation') for m in actual['members']):error='No active confirmation to reopen.'
            elif action=='repair' and not any(repair_eligible(m) for m in actual['members']):error='Not eligible for automatic repair.'
            elif action=='confirm':
                primary=actual['members'][0]
                if primary['affected_component']=='raw_price':
                    for other in all_findings(report):
                        if other['affected_component']=='adjusted_price' and other['ticker']==primary['ticker'] and other.get('overridable') and not (other.get('confirmation') and not other.get('inherited_from')):
                            dates=sorted(set(primary.get('session_dates',[])) & set(other.get('raw_missing_dates',[])))
                            if dates:item['inherited'].append(dict(issue_id=other['issue_id'],dates=dates))
        if error:item.update(status='REJECTED',error=error)
        result.append(item);duplicates.setdefault(gid,[]).append(item)
    for items in duplicates.values():
        if len(items)<2:continue
        if len({(i['action'],i['reason'],i['reference']) for i in items})>1:
            for item in items:item.update(status='REJECTED',error='Conflicting duplicate decisions for the same review group.')
        else:
            for item in items[1:]:
                if item['status']!='REJECTED':item.update(status='UNCHANGED',error='Identical duplicate row ignored.')
    # Reject conflicting actions on dependencies split across review groups too.
    for first in result:
        if first['status']!='VALID':continue
        touched=set(first['issue_ids'])|{x['issue_id'] for x in first['inherited']}
        for second in result:
            if second is first or second['status']!='VALID':continue
            if touched & set(second['issue_ids']):
                first.update(status='REJECTED',error='Overlapping review decisions; keep one action for this linked gap.')
                second.update(status='REJECTED',error='Overlapping review decisions; keep one action for this linked gap.')
    return result

def repair_eligible(row):
    return (row.get('resolution_status', 'UNRESOLVED') == 'UNRESOLVED'
            and row.get('affected_component') in ('raw_price','adjusted_price','adjustment_factor')
            and row.get('classification') in ('UNEXPLAINED_MISSING','INSUFFICIENT_EVIDENCE')
            and bool(row.get('session_dates')))

def preview_review(paths, report, encoded, filename='', progress=lambda **kw:None):
    with DATA_LOCK:
        snapshot,rows=parse_workbook(paths,encoded,progress)
        if snapshot['universe']!=report['universe']:raise ValueError('Workbook belongs to a different universe.')
        ledger_path=folder(paths)/'applied.json';applied=read_json(ledger_path) if ledger_path.exists() else {}
        progress(stage='VALIDATING',current=0,total=len(rows),message='Checking decisions against current findings.')
        decisions=evaluate(snapshot,rows,report,applied);preview_id=uuid.uuid4().hex
        progress(stage='VALIDATING',current=len(rows),total=len(rows),message='Decisions validated; saving preview.')
        write_json(folder(paths)/f'preview-{preview_id}.json',dict(snapshot=snapshot,rows=rows,universe=report['universe'],filename=filename))
        return dict(ok=True,preview_id=preview_id,filename=filename,rows_read=len(rows),decisions=decisions,counts=dict(Counter(d['status'] if d['status']!='VALID' else d['action'] for d in decisions)),inherited_count=sum(len(d['inherited']) for d in decisions if d['status']=='VALID'))


def load_preview(paths, report, preview_id):
    saved=read_json(folder(paths)/f'preview-{safe_id(preview_id)}.json')
    if saved['universe']!=report['universe']:raise ValueError('Preview belongs to another dataset.')
    ledger_path=folder(paths)/'applied.json';applied=read_json(ledger_path) if ledger_path.exists() else {}
    decisions=evaluate(saved['snapshot'],saved['rows'],report,applied)
    return dict(ok=True,preview_id=preview_id,filename=saved.get('filename',''),rows_read=len(saved['rows']),
                decisions=decisions,counts=dict(Counter(d['status'] if d['status']!='VALID' else d['action'] for d in decisions)),
                inherited_count=sum(len(d['inherited']) for d in decisions if d['status']=='VALID'))

def apply_review(paths, universe, preview_id, validate, repair, progress=lambda **kw:None):
    with DATA_LOCK:
        saved=read_json(folder(paths)/f'preview-{safe_id(preview_id)}.json')
        if saved['universe']!=universe:raise ValueError('Preview belongs to another dataset.')
        ledger_path=folder(paths)/'applied.json';ledger=read_json(ledger_path) if ledger_path.exists() else {}
        report=validate();decisions=evaluate(saved['snapshot'],saved['rows'],report,ledger)
        if not any(d['status']=='VALID' for d in decisions):
            raise ValueError('No valid actionable review decisions remain. Inspect the preview and select confirm, repair, or reopen before applying.')
        findings={r['issue_id']:r for r in all_findings(report)};repairs=[]
        for index,decision in enumerate(decisions,1):
            progress(stage='APPLYING',current=index-1,total=len(decisions),message=f'{decision["group_id"]}: {decision["action"]}')
            if decision['status']!='VALID':continue
            try:
                issue=findings[decision['primary_id']]
                if decision['action']=='confirm':
                    save_acknowledgement(paths,issue,'MANUALLY_ACCEPTED_GAP',decision['reason'],decision['reference'],findings=list(findings.values()),batch_id=preview_id,decision_id=decision['decision_id'],source_file=saved.get('filename'))
                elif decision['action']=='reopen':revoke_acknowledgement(paths,issue.get('reopen_issue_id',issue['issue_id']),batch_id=preview_id,decision_id=decision['decision_id'])
                elif decision['action']=='repair':repairs.append(decision);continue
                ledger[decision['decision_id']]=dict(batch_id=preview_id,action=decision['action'],timestamp=datetime.now(timezone.utc).isoformat());write_json(ledger_path,ledger)
                decision['status']='APPLIED'
            except Exception as exc:decision.update(status='REJECTED',error=str(exc))
        repair_result=None
        if repairs:
            ids=list(dict.fromkeys(i for d in repairs for i in d['issue_ids']))
            repair_result=repair(ids)
            for decision in repairs:
                decision['status']='APPLIED';ledger[decision['decision_id']]=dict(batch_id=preview_id,action='repair',committed=repair_result.get('committed'),
                    results=[r for r in repair_result.get('results',[]) if r['issue_id'] in decision['issue_ids']])
            write_json(ledger_path,ledger)
        progress(stage='REFRESHING_REPORT',current=len(decisions),total=len(decisions),message='Revalidating saved decisions and findings.')
        return dict(ok=True,batch_id=preview_id,decisions=decisions,repair=repair_result,report=validate())
