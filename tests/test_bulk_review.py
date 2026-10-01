"""Excel export/edit/preview/apply and repair integration, temporary data only."""
import base64
import copy
import hashlib
import io
import json
import unittest
from pathlib import Path
from unittest.mock import Mock, patch
import pandas as pd
from openpyxl import load_workbook
import test_completeness as fixtures
from momentumlab.issue_review import export_review, preview_review, apply_review, review_groups
from momentumlab.issue_repair import repair_issues
from momentumlab.data_integrity import load_acknowledgements, save_acknowledgement, clear_integrity_cache

class BulkReviewTests(unittest.TestCase):
    def setUp(self):
        self.f=fixtures.CompletenessTests();self.f.setUp();self.addCleanup(self.f.doCleanups)
        self.f.drop('raw_price','A',[16]);self.f.drop('adjusted_price','A',[16])
        self.paths=self.f.paths
    def export(self):return export_review(self.paths,self.f.report(mode='full'),{'A':'测试ETF'})
    def edit(self,export,action='confirm',reason='已核对，原因仍未证实；accepted gap',change=None):
        wb=load_workbook(io.BytesIO(base64.b64decode(export['workbook'])))
        sheet=wb['Review'];target=None
        for row in sheet.iter_rows(min_row=2):
            if row[5].value=='A' and 'raw_price' in row[8].value:
                row[16].value=action;row[17].value=reason;row[18].value='https://example.test/证据';target=row[0].row;break
        if change:change(sheet,target)
        stream=io.BytesIO();wb.save(stream);return base64.b64encode(stream.getvalue()).decode()
    def apply(self,preview):return apply_review(self.paths,'CSI300',preview['preview_id'],lambda:self.f.report(mode='full'),Mock())
    def test_roundtrip_chinese_group_inheritance_and_idempotency(self):
        exported=self.export();encoded=self.edit(exported)
        preview=preview_review(self.paths,self.f.report(),encoded)
        self.assertEqual(preview['counts']['confirm'],1);self.assertEqual(preview['inherited_count'],1)
        self.assertEqual(load_acknowledgements(self.paths),[]) # preview does not save decisions
        result=self.apply(preview);self.assertTrue(result['report']['research_allowed'])
        saved=load_acknowledgements(self.paths);self.assertEqual(len(saved),2)
        self.assertEqual(saved[0]['explanation'],'已核对，原因仍未证实；accepted gap');self.assertEqual(saved[1]['inherited_from'],saved[0]['confirmation_id'])
        self.assertEqual(saved[0]['import_batch_id'],preview['preview_id'])
        again=preview_review(self.paths,self.f.report(),encoded);self.assertEqual(again['counts']['ALREADY_APPLIED'],1)
        events=(self.f.root/'integrity_audit.jsonl').read_text().count('COMPLETENESS_CONFIRMED')
        with self.assertRaisesRegex(ValueError,'No valid actionable'):self.apply(again)
        self.assertEqual((self.f.root/'integrity_audit.jsonl').read_text().count('COMPLETENESS_CONFIRMED'),events)
        clear_integrity_cache();self.assertEqual(self.f.report()['status'],'WARNING')
    def test_upload_progress_uses_actual_workbook_row_count(self):
        exported=self.export();updates=[]
        result=preview_review(self.paths,self.f.report(),self.edit(exported),progress=lambda **values:updates.append(values))
        self.assertEqual(result['rows_read'],len(result['decisions']))
        self.assertTrue(any(x.get('stage')=='PARSING' and x.get('current')==result['rows_read'] for x in updates))
        self.assertTrue(any(x.get('stage')=='VALIDATING' and x.get('current')==result['rows_read'] for x in updates))
    def test_reordered_rows_and_protected_unknown_formula_duplicate_rejections(self):
        exported=self.export()
        variations=[lambda s,r:setattr(s.cell(r,6),'value','WRONG'),lambda s,r:setattr(s.cell(r,1),'value','UNKNOWN'),
                    lambda s,r:setattr(s.cell(r,18),'value','=1+1'),lambda s,r:s.append([c.value for c in s[r]][:16]+['reopen','different',''])]
        for change in variations:
            preview=preview_review(self.paths,self.f.report(),self.edit(exported,change=change))
            self.assertGreater(preview['counts'].get('REJECTED',0),0)
        encoded=self.edit(exported,change=lambda s,r:s.move_range(f'A{r}:S{r}',rows=100))
        self.assertEqual(preview_review(self.paths,self.f.report(),encoded)['counts']['confirm'],1)
    def test_reason_required_stale_evidence_and_apply_revalidation(self):
        exported=self.export()
        self.assertGreater(preview_review(self.paths,self.f.report(),self.edit(exported,reason=' '))['counts']['REJECTED'],0)
        preview=preview_review(self.paths,self.f.report(),self.edit(exported))
        self.f.life.loc[0,'list_date']=self.f.days[1];self.f.life.to_csv(self.f.root/'securities.csv',index=False)
        with self.assertRaisesRegex(ValueError,'No valid actionable'):self.apply(preview)
        self.assertEqual(load_acknowledgements(self.paths),[])
    def test_dataset_fingerprint_change_without_relevant_change_is_accepted(self):
        exported=self.export();report=self.f.report();report['manifest']['dataset_id']='new';report['manifest']['dataset_fingerprint']='new'
        self.assertEqual(preview_review(self.paths,report,self.edit(exported))['counts']['confirm'],1)
    def test_independent_adjusted_confirmation_preserved(self):
        report=self.f.report();child=next(r for r in report['completeness']['findings'] if r['affected_component']=='adjusted_price')
        original=save_acknowledgement(self.paths,child,'','Independent explanation')
        exported=self.export();result=self.apply(preview_review(self.paths,self.f.report(),self.edit(exported)))
        saved=load_acknowledgements(self.paths);independent=next(r for r in saved if r['confirmation_id']==original['confirmation_id'])
        self.assertTrue(independent['active']);self.assertEqual(independent['explanation'],'Independent explanation')
        actual=next(r for r in result['report']['completeness']['findings'] if r['affected_component']=='adjusted_price')
        self.assertEqual(actual['explanation'],'Independent explanation');self.assertIsNone(actual.get('inherited_from'))
    def test_bulk_reopen(self):
        self.apply(preview_review(self.paths,self.f.report(),self.edit(self.export())))
        result=self.apply(preview_review(self.paths,self.f.report(),self.edit(self.export(),action='reopen')))
        self.assertFalse(result['report']['research_allowed']);self.assertTrue(all(not r['active'] for r in load_acknowledgements(self.paths)))
    def test_full_export_not_representative_sample(self):
        report=self.f.report();report['completeness']['examples']=[]
        output=export_review(self.paths,report);wb=load_workbook(io.BytesIO(base64.b64decode(output['workbook'])))
        self.assertEqual(wb['Review'].max_row-1,len(review_groups(report)));self.assertGreater(output['findings'],0)

    def repair(self,provider):
        report=self.f.report();ids=[r['issue_id'] for r in report['completeness']['findings'] if r['severity']=='FAIL']
        return repair_issues(self.paths,'CSI300',ids,lambda:self.f.report(mode='full'),provider,self.f.ns['build_adjusted_prices'])
    def hashes(self):return {k:hashlib.sha256(Path(v).read_bytes()).hexdigest() for k,v in self.paths.items() if k!='root'}
    def test_empty_provider_is_not_repaired(self):
        before=self.hashes();result=self.repair(lambda *args:pd.DataFrame())
        self.assertFalse(result['committed']);self.assertTrue(all(r['status']=='Still missing' for r in result['results']));self.assertEqual(before,self.hashes())
    def source(self,api,params,fields):
        frame=self.f.raw if api=='daily' else self.f.fac
        return frame[(frame.ts_code==params['ts_code']) & frame.trade_date.eq(pd.Timestamp(params['start_date']))].copy()
    def test_repair_success_and_publication_failure_preserves_data(self):
        before=self.hashes()
        with patch('momentumlab.issue_repair.publish_dataset',side_effect=OSError('simulated commit failure')):
            result=self.repair(self.source)
        self.assertFalse(result['committed']);self.assertEqual(before,self.hashes())
        result=self.repair(self.source);self.assertTrue(result['committed']);self.assertTrue(all(r['status']=='Repaired and revalidated' for r in result['results']))
    def test_partial_provider_error_publishes_only_validated_improvement(self):
        self.f.drop('raw_price','B',[16]);self.f.drop('adjusted_price','B',[16]);before=self.hashes()
        def provider(api,params,fields):
            if params['ts_code']=='B':raise RuntimeError('Provider unavailable for B')
            return self.source(api,params,fields)
        result=self.repair(provider);self.assertTrue(result['committed'])
        self.assertNotEqual(before,self.hashes())
        self.assertIn('Provider error',[r['status'] for r in result['results']])
        self.assertIn('Repaired and revalidated',[r['status'] for r in result['results']])
        self.assertFalse(self.f.report(mode='full')['research_allowed'])

    def test_confirmed_findings_do_not_trigger_provider_repair(self):
        report=self.f.report(mode='full')
        raw=next(r for r in report['completeness']['findings'] if r['affected_component']=='raw_price')
        save_acknowledgement(self.paths,raw,'MANUALLY_ACCEPTED_GAP','Reviewed known gap',findings=report['completeness']['findings'])
        calls=[]
        result=self.repair(lambda *args:calls.append(args) or pd.DataFrame())
        self.assertEqual(calls,[])
        self.assertFalse(result['committed'])
        self.assertTrue(all(row['status']=='Not eligible for automatic repair' for row in result['results']))

    def test_missing_factor_does_not_redownload_present_raw_price(self):
        self.f.drop('adj_factor','B',[16])
        calls=[]
        def provider(api,params,fields):
            calls.append(api)
            return self.source(api,params,fields)
        report=self.f.report(mode='full')
        ids=[r['issue_id'] for r in report['completeness']['findings'] if r['ticker']=='B' and r['affected_component']=='adjustment_factor']
        result=repair_issues(self.paths,'CSI300',ids,lambda:self.f.report(mode='full'),provider,self.f.ns['build_adjusted_prices'])
        self.assertEqual(calls,['adj_factor'])
        self.assertTrue(result['committed'],result.get('error'))

    def test_etf_date_batches_and_exact_filtering(self):
        from momentumlab.data_integrity import validate_dataset
        self.f.drop('raw_price','B',[16]);self.f.drop('adjusted_price','B',[16])
        weights=pd.read_csv(self.paths['weights_csv']);weights['market_cap_rmb']=300000000;weights.to_csv(self.paths['weights_csv'],index=False)
        pd.DataFrame([dict(ts_code=t,name='Test ETF',fund_type='Equity') for t in ['A','B']]).to_csv(self.f.root/'fund_classification.csv',index=False)
        validate=lambda:validate_dataset(self.paths,universe_id='ETF_250M',mode='full',lookback=10)
        report=validate();ids=[r['issue_id'] for r in report['completeness']['findings'] if r['severity']=='FAIL'];calls=[]
        def provider(api,params,fields):
            calls.append((api,params));frame=self.f.raw.copy() if api=='fund_daily' else self.f.fac.copy()
            return pd.concat([frame,frame.assign(ts_code='UNSELECTED')]) # Out-of-range/ticker rows must be ignored.
        result=repair_issues(self.paths,'ETF_250M',ids,validate,provider,self.f.ns['build_adjusted_prices'])
        self.assertTrue(result['committed'],result.get('error'));self.assertEqual(len(calls),1)
        self.assertEqual(calls[0][0],'fund_daily')
        self.assertEqual(calls[0][1],{'trade_date':self.f.days[16].strftime('%Y%m%d')})
        raw=pd.read_csv(self.paths['raw_price_csv']);self.assertNotIn('UNSELECTED',set(raw.ts_code));self.assertFalse(raw.duplicated(['ts_code','trade_date']).any())

if __name__=='__main__':unittest.main()
